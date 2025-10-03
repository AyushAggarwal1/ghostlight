from __future__ import annotations

import io
import os
from typing import Optional, List

import chardet  # type: ignore

try:
    from pdfminer.high_level import extract_text as pdf_extract_text  # type: ignore
except Exception:  # pragma: no cover
    pdf_extract_text = None

try:
    import pdfplumber  # type: ignore
except Exception:  # pragma: no cover
    pdfplumber = None

try:
    import docx  # type: ignore
except Exception:  # pragma: no cover
    docx = None

try:
    from PIL import Image  # type: ignore
    import pytesseract  # type: ignore
except Exception:  # pragma: no cover
    Image = None  # type: ignore
    pytesseract = None  # type: ignore

import zipfile
import tarfile
try:
    import py7zr  # type: ignore
except Exception:  # pragma: no cover
    py7zr = None


def read_text_bytes(path: str, sample_bytes: int) -> str:
    with open(path, "rb") as fh:
        raw = fh.read(sample_bytes)
    detection = chardet.detect(raw)
    encoding = detection.get("encoding") or "utf-8"
    try:
        return raw.decode(encoding, errors="ignore")
    except Exception:
        return raw.decode("utf-8", errors="ignore")


def extract_text_from_pdf(path: str, max_chars: int = 5000) -> Optional[str]:
    # Prefer pdfplumber for layout-preserved extraction; fallback to pdfminer
    try:
        if pdfplumber is not None:
            with pdfplumber.open(path) as pdf:
                pages = [p.extract_text(x_tolerance=1, y_tolerance=1) or "" for p in pdf.pages[:10]]
                txt = "\n\n".join(pages)
                if txt.strip():
                    return txt[:max_chars]
    except Exception:
        pass
    if pdf_extract_text is not None:
        try:
            text = pdf_extract_text(path) or ""
            return text[:max_chars]
        except Exception:
            return None
    return None


def extract_text_from_docx(path: str, max_chars: int = 5000) -> Optional[str]:
    if docx is None:
        return None
    try:
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        return text[:max_chars]
    except Exception:
        return None


def extract_text_from_excel(path: str, max_chars: int = 5000) -> Optional[str]:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text_parts.append(" ".join(str(cell) if cell is not None else "" for cell in row))
                if len(" ".join(text_parts)) > max_chars:
                    break
        return " ".join(text_parts)[:max_chars]
    except Exception:
        return None


def extract_text_from_csv(path: str, max_chars: int = 5000) -> Optional[str]:
    try:
        import pandas as pd  # type: ignore
        df = pd.read_csv(path, nrows=100)
        return df.to_string()[:max_chars]
    except Exception:
        return None


def extract_text_from_image(path: str, max_chars: int = 5000) -> Optional[str]:
    if Image is None or pytesseract is None:
        return None
    try:
        img = Image.open(path)
        img = img.convert("L")
        text = pytesseract.image_to_string(img) or ""
        return text[:max_chars]
    except Exception:
        return None


def extract_texts_from_archive(path: str, sample_bytes: int = 2048, max_files: int = 10) -> Optional[str]:
    texts: List[str] = []
    try:
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as z:
                for info in z.infolist()[:max_files]:
                    if info.is_dir():
                        continue
                    with z.open(info) as fh:
                        data = fh.read(sample_bytes)
                        texts.append(data.decode("utf-8", errors="ignore"))
            return "\n".join(texts)
        if tarfile.is_tarfile(path):
            with tarfile.open(path, "r:*") as t:
                count = 0
                for member in t.getmembers():
                    if member.isreg():
                        f = t.extractfile(member)
                        if f is None:
                            continue
                        data = f.read(sample_bytes)
                        texts.append(data.decode("utf-8", errors="ignore"))
                        count += 1
                        if count >= max_files:
                            break
            return "\n".join(texts)
        if py7zr is not None and path.lower().endswith(".7z"):
            with py7zr.SevenZipFile(path, mode="r") as z:
                for name, bio in list(z.readall().items())[:max_files]:
                    data = bio.read(sample_bytes)
                    texts.append(data.decode("utf-8", errors="ignore"))
            return "\n".join(texts)
    except Exception:
        return None
    return None


def extract_text_from_file(path: str, sample_bytes: int = 2048) -> str:
    _, ext = os.path.splitext(path)
    ext = ext.lower()
    if ext == ".pdf":
        pdf_text = extract_text_from_pdf(path)
        if pdf_text:
            return pdf_text
    if ext in {".docx"}:
        doc_text = extract_text_from_docx(path)
        if doc_text:
            return doc_text
    if ext in {".xlsx", ".xls"}:
        excel_text = extract_text_from_excel(path)
        if excel_text:
            return excel_text
    if ext == ".csv":
        csv_text = extract_text_from_csv(path)
        if csv_text:
            return csv_text
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"}:
        ocr_text = extract_text_from_image(path)
        if ocr_text:
            return ocr_text
    if ext in {".zip", ".tar", ".gz", ".tgz", ".bz2", ".7z"}:
        arc_text = extract_texts_from_archive(path, sample_bytes)
        if arc_text:
            return arc_text
    # Default to bytes read + encoding detection
    return read_text_bytes(path, sample_bytes)


