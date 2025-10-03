from __future__ import annotations

import io
import os
from typing import Optional

import chardet  # type: ignore

try:
    from pdfminer.high_level import extract_text as pdf_extract_text  # type: ignore
except Exception:  # pragma: no cover
    pdf_extract_text = None

try:
    import docx  # type: ignore
except Exception:  # pragma: no cover
    docx = None


def read_text_bytes(path: str, sample_bytes: int) -> str:
    with open(path, "rb") as fh:
        raw = fh.read(sample_bytes)
    detection = chardet.detect(raw)
    encoding = detection.get("encoding") or "utf-8"
    try:
        return raw.decode(encoding, errors="ignore")
    except Exception:
        return raw.decode("utf-8", errors="ignore")


def extract_text_from_pdf(path: str, max_chars: int = 5000) -> Optional[str]:
    if pdf_extract_text is None:
        return None
    try:
        text = pdf_extract_text(path) or ""
        return text[:max_chars]
    except Exception:
        return None


def extract_text_from_docx(path: str, max_chars: int = 5000) -> Optional[str]:
    if docx is None:
        return None
    try:
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        return text[:max_chars]
    except Exception:
        return None


def extract_text_from_excel(path: str, max_chars: int = 5000) -> Optional[str]:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text_parts.append(" ".join(str(cell) if cell is not None else "" for cell in row))
                if len(" ".join(text_parts)) > max_chars:
                    break
        return " ".join(text_parts)[:max_chars]
    except Exception:
        return None


def extract_text_from_csv(path: str, max_chars: int = 5000) -> Optional[str]:
    try:
        import pandas as pd  # type: ignore
        df = pd.read_csv(path, nrows=100)
        return df.to_string()[:max_chars]
    except Exception:
        return None


def extract_text_from_file(path: str, sample_bytes: int = 2048) -> str:
    _, ext = os.path.splitext(path)
    ext = ext.lower()
    if ext == ".pdf":
        pdf_text = extract_text_from_pdf(path)
        if pdf_text:
            return pdf_text
    if ext in {".docx"}:
        doc_text = extract_text_from_docx(path)
        if doc_text:
            return doc_text
    if ext in {".xlsx", ".xls"}:
        excel_text = extract_text_from_excel(path)
        if excel_text:
            return excel_text
    if ext == ".csv":
        csv_text = extract_text_from_csv(path)
        if csv_text:
            return csv_text
    # Default to bytes read + encoding detection
    return read_text_bytes(path, sample_bytes)


